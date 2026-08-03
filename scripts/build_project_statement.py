#!/usr/bin/env python3
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_EN = ROOT / "docs" / "MOTION_INSTRUMENT_Project_Statement_Draft_EN.docx"
OUTPUT_ZH = ROOT / "docs" / "MOTION_INSTRUMENT_Project_Statement_Draft_ZH.docx"
FONT_NAME = "Calibri"

GREEN = RGBColor(31, 74, 61)
GOLD = RGBColor(224, 181, 52)
INK = RGBColor(30, 44, 39)
MUTED = RGBColor(104, 128, 118)
PALE = "EAF1ED"


def set_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_NAME)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def shade_paragraph(paragraph, fill=PALE):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_body(doc, text, after=8):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.333
    set_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    if level == 1:
        size, before, after, color = 16, 18, 10, GREEN
    elif level == 2:
        size, before, after, color = 13, 12, 6, GREEN
    else:
        size, before, after, color = 12, 8, 4, INK
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    set_font(paragraph.add_run(text), size=size, color=color, bold=True)
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    set_font(paragraph.add_run(text), size=9, color=MUTED, italic=True)


def add_image(doc, relative_path, width=6.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(ROOT / relative_path), width=Inches(width))


def add_module(doc, title, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.15
    set_font(paragraph.add_run(f"{title}. "), size=10.5, color=GREEN, bold=True)
    set_font(paragraph.add_run(text), size=10.5)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_font(run, size=9, color=MUTED)


def configure_document(doc, header_text):
    section = doc.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run(header_text), size=8.5, color=MUTED, bold=True)
    add_page_number(section.footer.paragraphs[0])


def build_english():
    global FONT_NAME
    FONT_NAME = "Calibri"
    doc = Document()
    configure_document(doc, "MOTION INSTRUMENT  |  PROJECT STATEMENT DRAFT")

    # Page 1 — editorial cover and project statement.
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(8)
    set_font(kicker.add_run("GESTURE-CONTROLLED LIVE SOUND PERFORMANCE SYSTEM"), size=9.5, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    set_font(title.add_run("MOTION INSTRUMENT"), size=27, color=GREEN, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_font(subtitle.add_run("A Max/MSP instrument for spatial hand roles, continuous control and live transformation"), size=12.5, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(14)
    set_font(meta.add_run(f"Project statement draft  |  {date.today().strftime('%d %B %Y')}"), size=9.5, color=MUTED, italic=True)

    add_image(doc, "assets/ui/concepts/perform_ui_concept_v7_final_cleanup.png", 6.5)
    add_caption(doc, "Perform-page design: sources, gesture feedback, serial effects, recording and master monitoring in one view.")
    add_heading(doc, "Project Statement", 1)
    add_body(doc, "Motion Instrument is a live sound performance system that treats bodily position as part of an instrument rather than as a remote control layered on top of one. Built in Max/MSP, it combines microphone, file and granular sources with camera-based hand tracking, four serial audio effects, recording and a performance-oriented interface.")
    add_body(doc, "The project asks how gesture mapping can remain readable during performance. Instead of trusting left/right hand labels, which can swap during tracking, the camera image is divided into stable spatial roles. One zone adjusts continuous parameters through X position, Y position and pinch; the other selects an effect with finger counts 1–5, while a fist freezes change. The performer can therefore see, hear and understand the current control state without leaving the main interface.", after=0)

    # Page 2 — motivation and interaction.
    page_heading = add_heading(doc, "1. Artistic Motivation", 1)
    page_heading.paragraph_format.page_break_before = True
    add_body(doc, "Electronic performance often separates expressive motion from the technical controls that shape sound. Mouse-driven interfaces are precise, but they pull attention toward a screen and reduce the visibility of cause and effect for an audience. Motion Instrument explores a more physical relationship: a movement should create an audible change, while the interface explains what the system believes the movement means.")
    add_body(doc, "The goal is not to replace tactile controllers. It is to create a complementary performance layer for gestures that benefit from scale, distance and visible intention. This led to a design that favours a small number of legible mappings, explicit feedback and a reliable hold state over a large vocabulary of symbolic gestures.")
    add_heading(doc, "2. Interaction Design", 1)
    add_image(doc, "assets/ui/concepts/gesture_detail_layout_v3_live_feedback.png", 6.35)
    add_caption(doc, "Gesture page: spatial work zones, target selection and two rows of numerical mapping feedback.")
    add_module(doc, "Parameter zone", "The active hand produces normalized X, Y and pinch values. The interface displays both the raw coordinates and the effect parameters that receive them.")
    add_module(doc, "Gesture gate", "Finger counts 1–4 select Vocoder, Bitcrusher, Multiband Filter and Feedback Delay; 5 targets all four effects. A fist enters Hold and prevents further mapping changes.")
    add_module(doc, "Readable feedback", "The recognized value 0–5 is shown immediately and the Control Target changes with the confirmed gesture, allowing the performer to correct a misread before it becomes musically disruptive.")

    # Page 3 — implementation.
    page_heading = add_heading(doc, "3. Technical Realisation", 1)
    page_heading.paragraph_format.page_break_before = True
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_before = Pt(0)
    flow.paragraph_format.space_after = Pt(10)
    flow.paragraph_format.left_indent = Inches(0.18)
    flow.paragraph_format.right_indent = Inches(0.18)
    shade_paragraph(flow)
    set_font(flow.add_run("MIC / FILE / GRANULAR  →  SOURCE MIXER  →  VOCODER  →  BITCRUSHER  →  MULTIBAND  →  FEEDBACK DELAY  →  RECORD / MONITOR"), size=9.5, color=GREEN, bold=True)
    add_body(doc, "MediaPipe hand tracking runs inside Max through jweb. The bridge sends landmark and gesture data to a shared mapping state machine, which assigns hands by screen position rather than handedness labels. Ambiguous occupancy is rejected, and a short stability period is required before a zone becomes active.")
    add_body(doc, "Camera refresh rate directly affects how often new control values arrive. To avoid a faster camera producing a more aggressive response, numerical smoothing uses elapsed timestamps with a 100 ms time constant instead of a fixed per-frame step. Gesture confirmation is also time based: number selection requires 300 ms, while the fist Hold requires 100 ms. Frame-count fallbacks remain only for inputs without timestamps.")
    add_image(doc, "assets/ui/concepts/fx_detail_layout_v5_source_style_actual_ratio.png", 6.35)
    add_caption(doc, "FX page: one full effect editor at a time, with four persistent output controls for the serial stages.")
    add_body(doc, "The audio path is deliberately serial. Vocoder and Bitcrusher establish the transformed material, Multiband Filter reshapes its spectrum, and Feedback Delay adds the final temporal layer. Click-free 20 ms bypass ramps keep the dry signal moving through disabled stages. The FX page mirrors each module's existing Output Gain rather than creating a second gain state.", after=0)

    # Page 4 — sonic intent, reflection and current stage.
    page_heading = add_heading(doc, "4. Sound Design Strategy", 1)
    page_heading.paragraph_format.page_break_before = True
    add_module(doc, "Granular source", "Creates sustained or fragmented material from a microphone capture or loaded sample. Wet, density, size and pitch are exposed as performance macros.")
    add_module(doc, "Vocoder", "Introduces a spectral voice layer and hybrid carrier character. Brightness, carrier tone, noise mix, dry/wet and spectral smoothing support both intelligible and synthetic results.")
    add_module(doc, "Bitcrusher", "Adds clearly perceptible digital degradation through sample-rate reduction, bit quantisation and drive. Its strong contrast makes gesture changes easy to hear.")
    add_module(doc, "Multiband Filter", "Divides the sound into low, mid and high regions for spectral focus, contrast and spread. It functions as the middle shaping stage rather than as a parallel return.")
    add_module(doc, "Feedback Delay", "Finishes the chain with delay time, feedback and stereo offset. Placing it last lets previous transformations generate a coherent echo tail.")
    add_heading(doc, "5. Reflection and Current Stage", 1)
    add_body(doc, "The main design lesson is that recognition accuracy alone does not create a playable interface. The performer also needs a stable role model, time-consistent response and visible confirmation. Spatial work zones, Hold and numerical feedback emerged from repeated camera testing and now define the identity of the instrument.")
    add_image(doc, "assets/ui/concepts/master_detail_layout_v4_record_monitor.png", 4.8)
    add_caption(doc, "Master page: non-destructive record, review and export alongside final monitoring.")
    add_body(doc, "The current Max 9 prototype has completed its structural and static test pass, including all active subpatch layouts. Before public release, the remaining work is a final Max runtime review of audio, camera behaviour, Console output and export on the presentation computer, followed by documentation cleanup and packaging. The project remains a performance prototype rather than a standalone plug-in.", after=0)

    OUTPUT_EN.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_EN)
    print(OUTPUT_EN)


def build_chinese():
    global FONT_NAME
    FONT_NAME = "Heiti SC"
    doc = Document()
    configure_document(doc, "MOTION INSTRUMENT  |  项目阐述初版")

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(10)
    kicker.paragraph_format.space_after = Pt(8)
    set_font(kicker.add_run("手势控制现场声音表演系统"), size=9.5, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    set_font(title.add_run("MOTION INSTRUMENT"), size=27, color=GREEN, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(10)
    set_font(subtitle.add_run("一件基于空间手势角色、连续控制与现场声音变换的 Max/MSP 乐器"), size=12.5, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(14)
    today = date.today()
    set_font(meta.add_run(f"项目阐述初版  |  {today.year}年{today.month}月{today.day}日"), size=9.5, color=MUTED, italic=True)

    add_image(doc, "assets/ui/concepts/perform_ui_concept_v7_final_cleanup.png", 6.5)
    add_caption(doc, "Perform 页面设计：在同一视图中呈现音源、手势反馈、串联效果器、录制与主监听。")
    add_heading(doc, "项目阐述", 1)
    add_body(doc, "Motion Instrument 是一套现场声音表演系统，它把身体位置视为乐器的一部分，而不是附加在乐器之外的遥控方式。系统以 Max/MSP 构建，将麦克风、文件与粒子合成音源，与摄像头手势追踪、四个串联音频效果器、录音导出和面向表演的界面整合在一起。")
    add_body(doc, "项目关注的核心问题，是如何让手势映射在表演过程中始终清晰可读。系统不依赖容易在追踪中互换的左右手标签，而是把摄像头画面划分为稳定的空间角色：一个区域通过 X、Y 位置和捏合距离连续调整参数；另一个区域通过 1–5 的手指数选择效果器，握拳则冻结变化。表演者无需离开主界面，就能看见、听见并理解当前控制状态。", after=0)

    page_heading = add_heading(doc, "1. 艺术动机", 1)
    page_heading.paragraph_format.page_break_before = True
    add_body(doc, "电子音乐表演常常把具有表达性的身体动作，与真正塑造声音的技术控制分离开来。鼠标界面虽然精确，却会把注意力拉回屏幕，也让观众难以理解动作与声音之间的因果关系。Motion Instrument 尝试建立一种更具身体性的关系：动作应当产生可以听见的变化，同时界面需要解释系统如何理解这个动作。")
    add_body(doc, "本项目并不试图取代实体控制器，而是为适合通过幅度、距离和可见意图表达的动作增加一层表演控制。因此，设计更重视少量且清晰的映射、明确的识别反馈和可靠的 Hold 状态，而不是追求大量象征性手势。")
    add_heading(doc, "2. 交互设计", 1)
    add_image(doc, "assets/ui/concepts/gesture_detail_layout_v3_live_feedback.png", 6.35)
    add_caption(doc, "Gesture 页面：空间工作区、目标选择，以及两排数值化映射反馈。")
    add_module(doc, "参数工作区", "进入该区域的手输出归一化 X、Y 与 Pinch 数值。界面同时显示输入坐标和实际接收这些数值的效果器参数。")
    add_module(doc, "手势选择区", "手指数 1–4 分别选择 Vocoder、Bitcrusher、Multiband Filter 与 Feedback Delay；5 同时控制四个效果器；握拳进入 Hold，停止映射变化。")
    add_module(doc, "可读反馈", "识别到的 0–5 会即时显示，Control Target 在手势确认后同步切换，使表演者能在错误识别影响声音之前及时修正。")

    page_heading = add_heading(doc, "3. 技术实现", 1)
    page_heading.paragraph_format.page_break_before = True
    flow = doc.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    flow.paragraph_format.space_before = Pt(0)
    flow.paragraph_format.space_after = Pt(10)
    flow.paragraph_format.left_indent = Inches(0.18)
    flow.paragraph_format.right_indent = Inches(0.18)
    shade_paragraph(flow)
    set_font(flow.add_run("麦克风 / 文件 / 粒子音源  →  SOURCE MIXER  →  VOCODER  →  BITCRUSHER  →  MULTIBAND  →  FEEDBACK DELAY  →  录制 / 监听"), size=9.5, color=GREEN, bold=True)
    add_body(doc, "MediaPipe 手势追踪通过 jweb 运行在 Max 内部。桥接层把手部关键点和手势数据送入共享映射状态机，并依据手在画面中的位置分配角色，而不是使用左右手标签。当同一区域出现多只手时，系统会拒绝模糊状态；手进入区域后还需经过短暂稳定期才会激活。")
    add_body(doc, "摄像头刷新率会直接改变控制数据的到达频率。为了避免高帧率摄像头产生更激烈的参数响应，数值平滑使用时间戳计算，并采用 100 ms 时间常数，而不是固定的逐帧步长。手势确认同样基于时间：数字选择需保持 300 ms，握拳 Hold 需保持 100 ms；只有缺少时间戳时才使用帧数作为后备。")
    add_image(doc, "assets/ui/concepts/fx_detail_layout_v5_source_style_actual_ratio.png", 6.35)
    add_caption(doc, "FX 页面：一次完整显示一个效果器，同时固定呈现四个串联阶段的输出控制。")
    add_body(doc, "音频路径采用串联结构。Vocoder 与 Bitcrusher 建立主要变换材质，Multiband Filter 重塑频谱，Feedback Delay 在最后增加时间层。20 ms 无点击旁通斜坡保证关闭某个阶段后干声仍能继续传递。FX 页面镜像各模块已有的 Output Gain，而不是建立第二套增益状态。", after=0)

    page_heading = add_heading(doc, "4. 声音设计策略", 1)
    page_heading.paragraph_format.page_break_before = True
    add_module(doc, "粒子音源", "从麦克风录音或加载的采样生成持续或碎片化声音，并以 Wet、Density、Size 与 Pitch 作为表演宏参数。")
    add_module(doc, "Vocoder", "引入频谱人声层与混合载波质感。Brightness、Carrier Tone、Noise Mix、Dry/Wet 和 Spectral Smooth 可在可辨识人声与合成音色之间变化。")
    add_module(doc, "Bitcrusher", "通过采样率降低、位深量化与 Drive 产生明显的数字退化，其强烈对比使手势变化更容易被听见。")
    add_module(doc, "Multiband Filter", "把声音分为低、中、高三个频段，用于频谱聚焦、对比与扩展；它是串联链中部的塑形阶段，而不是并联 Return。")
    add_module(doc, "Feedback Delay", "以 Delay Time、Feedback 与 Stereo Offset 完成效果链。放置在最后，可让之前的声音变换形成连贯的回声尾音。")
    add_heading(doc, "5. 反思与当前阶段", 1)
    add_body(doc, "项目最重要的设计经验是：识别准确率本身并不能形成可演奏的界面。表演者还需要稳定的角色逻辑、不受帧率影响的响应，以及明确的视觉确认。空间工作区、Hold 与数值反馈都来自反复的摄像头测试，并最终成为这件乐器的核心特征。")
    add_image(doc, "assets/ui/concepts/master_detail_layout_v4_record_monitor.png", 4.8)
    add_caption(doc, "Master 页面：无损录制、试听与导出，以及最终监听控制。")
    add_body(doc, "当前 Max 9 原型已经完成结构与静态测试，包括全部活动 subpatch 的布局检查。公开发布之前，仍需在展示电脑上完成音频、摄像头行为、Console 输出与导出的最终 Max 运行检查，随后整理发布文档并打包。本项目目前仍是一套表演原型，而不是独立插件。", after=0)

    OUTPUT_ZH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_ZH)
    print(OUTPUT_ZH)


if __name__ == "__main__":
    build_english()
    build_chinese()
